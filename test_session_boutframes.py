"""
Tests for session-aware boutframes lookup.

Why this exists: a subject ID is one *recording* (`CAB01_Alcohol`), and a
repeated-measures project has several per animal.  Boutframes used to be looked
up as ``read_excel(file, sheet_name=subject_id)`` -- one sheet per recording and
nothing else.  A workbook that instead keys sheets by ANIMAL and tags each row
with the session it was scored in (what tools/combine_sessions.py writes) has to
resolve to the same bouts, and a single-session project must be unaffected.

The methods under test are Tk-free; they are bound to a minimal stub the same
way test_animal_identity.py does it.
"""

import importlib.util
import os
import sys

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

_GUI_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fp_analysis_gui.py')


def _load_gui_module():
    spec = importlib.util.spec_from_file_location('fp_analysis_gui_bfsession', _GUI_PATH)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)   # safe: the GUI only starts under __main__
    return module


_M = _load_gui_module()
_G = _M.FPAnalysisGUI

UNDERSCORE = _M.SESSION_PATTERN_PRESETS['Subject_Session  (2F_pre, 2F_post)']
NONE = _M.SESSION_PATTERN_PRESETS['No sessions      (each ID is one animal)']


class _Bouts:
    """Stub exposing the boutframes helpers; they touch only params/processed_data."""
    BOUTFRAMES_SESSION_COLUMNS = _G.BOUTFRAMES_SESSION_COLUMNS
    _session_pattern = _G._session_pattern
    _boutframes_sheet_index = _G._boutframes_sheet_index
    _session_column = _G._session_column
    resolve_boutframes_sheet = _G.resolve_boutframes_sheet
    read_boutframes_sheet = _G.read_boutframes_sheet
    filter_boutframes_rows = _G.filter_boutframes_rows
    iter_boutframes_recordings = _G.iter_boutframes_recordings
    boutframes_behavior_columns = _G.boutframes_behavior_columns
    boutframes_recording_for_sheet = _G.boutframes_recording_for_sheet
    _parse_boutframes_dataframe = _G._parse_boutframes_dataframe
    # staticmethod() re-applied: assigning the plain function into a class body
    # would rebind it as an instance method and shift its arguments.
    _read_sheet_columns = staticmethod(_G._read_sheet_columns)
    _read_sheet_columns_by_session = _G._read_sheet_columns_by_session
    _rewrite_sheet_columns_by_session = _G._rewrite_sheet_columns_by_session

    def __init__(self, subjects=(), pattern=UNDERSCORE):
        self.params = {'session_pattern': pattern, 'animal_overrides': {}}
        self.processed_data = {s: {} for s in subjects}
        self.log_message = lambda *a, **k: None


# ---------------------------------------------------------------------------
# Fixtures: the two workbook layouts
# ---------------------------------------------------------------------------

@pytest.fixture
def per_animal(tmp_path):
    """One sheet per animal, rows tagged with the session they were scored in."""
    path = tmp_path / 'boutframes_alldays.xlsx'
    cab01 = pd.DataFrame({
        'Session':         ['Alcohol', 'Alcohol', 'Fentanyl', 'WaterSucrose'],
        'Alcohol__start':  [156, 525, None, None],
        'Alcohol__end':    [327, 612, None, None],
        'Fentanyl__start': [None, None, 2567, None],
        'Fentanyl__end':   [None, None, 2700, None],
        'Water__start':    [None, None, None, 1097],
        'Water__end':      [None, None, None, 1168],
    })
    cab03 = pd.DataFrame({
        'Session':        ['Alcohol'],
        'Alcohol__start': [900],
        'Alcohol__end':   [950],
    })
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        cab01.to_excel(w, sheet_name='CAB01', index=False)
        cab03.to_excel(w, sheet_name='CAB03', index=False)
    return str(path)


@pytest.fixture
def per_recording(tmp_path):
    """The pre-existing layout: one sheet named for each recording."""
    path = tmp_path / 'boutframes.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        pd.DataFrame({'Alcohol__start': [156], 'Alcohol__end': [327]}).to_excel(
            w, sheet_name='CAB01_Alcohol', index=False)
        pd.DataFrame({'Fentanyl__start': [2567], 'Fentanyl__end': [2700]}).to_excel(
            w, sheet_name='CAB01_Fentanyl', index=False)
    return str(path)


@pytest.fixture
def single_session(tmp_path):
    """A project with no sessions at all -- must behave exactly as before."""
    path = tmp_path / 'boutframes.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        pd.DataFrame({'Drink__start': [10, 20], 'Drink__end': [15, 25]}).to_excel(
            w, sheet_name='DG01', index=False)
    return str(path)


# ---------------------------------------------------------------------------
# Resolution
# ---------------------------------------------------------------------------

def test_a_recording_resolves_to_its_animals_sheet(per_animal):
    b = _Bouts(['CAB01_Alcohol'])
    assert b.resolve_boutframes_sheet(per_animal, 'CAB01_Alcohol') == ('CAB01', 'Alcohol')


def test_a_sheet_named_for_the_recording_wins_over_the_animals(tmp_path):
    """An exact per-recording sheet is unambiguous, so it must not be overridden."""
    path = tmp_path / 'bf.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        pd.DataFrame({'Session': ['Alcohol'], 'A__start': [1]}).to_excel(
            w, sheet_name='CAB01', index=False)
        pd.DataFrame({'A__start': [999]}).to_excel(
            w, sheet_name='CAB01_Alcohol', index=False)
    b = _Bouts(['CAB01_Alcohol'])
    assert b.resolve_boutframes_sheet(str(path), 'CAB01_Alcohol') == ('CAB01_Alcohol', '')


def test_sheet_names_match_case_insensitively(per_animal):
    b = _Bouts(['cab01_Alcohol'])
    sheet, session = b.resolve_boutframes_sheet(per_animal, 'cab01_Alcohol')
    assert sheet == 'CAB01' and session == 'Alcohol'


def test_resolution_works_before_anything_is_processed(per_animal):
    """extract_bouts runs from inside process_fp_data, BEFORE the subject lands
    in processed_data -- for the first subject the project is still empty.  So
    resolution must read the ID itself, not the project-level animal map."""
    b = _Bouts()          # nothing processed
    assert b.resolve_boutframes_sheet(per_animal, 'CAB01_Fentanyl') == ('CAB01', 'Fentanyl')
    assert list(b.read_boutframes_sheet(per_animal, 'CAB01_Fentanyl')['Fentanyl__start']) == [2567]


def test_an_unknown_recording_resolves_to_nothing(per_animal):
    b = _Bouts(['CAB99_Alcohol'])
    assert b.resolve_boutframes_sheet(per_animal, 'CAB99_Alcohol') == (None, '')


def test_reading_an_unresolvable_recording_raises(per_animal):
    b = _Bouts(['CAB99_Alcohol'])
    with pytest.raises(Exception):
        b.read_boutframes_sheet(per_animal, 'CAB99_Alcohol')


def test_an_animal_sheet_is_not_used_when_ids_encode_no_session(per_animal):
    """With session parsing off, `CAB01_Alcohol` is a whole animal ID and must
    not silently borrow CAB01's bouts."""
    b = _Bouts(['CAB01_Alcohol'], pattern=NONE)
    assert b.resolve_boutframes_sheet(per_animal, 'CAB01_Alcohol') == (None, '')


# ---------------------------------------------------------------------------
# Row filtering -- the part that keeps two days from bleeding into each other
# ---------------------------------------------------------------------------

def test_each_recording_gets_only_its_own_sessions_rows(per_animal):
    b = _Bouts(['CAB01_Alcohol', 'CAB01_Fentanyl', 'CAB01_WaterSucrose'])

    alc = b.read_boutframes_sheet(per_animal, 'CAB01_Alcohol')
    assert list(alc['Alcohol__start']) == [156, 525]
    assert alc['Fentanyl__start'].isna().all()

    fen = b.read_boutframes_sheet(per_animal, 'CAB01_Fentanyl')
    assert list(fen['Fentanyl__start']) == [2567]
    assert fen['Alcohol__start'].isna().all()


def test_the_session_column_is_dropped_so_it_is_never_a_behavior(per_animal):
    b = _Bouts(['CAB01_Alcohol'])
    df = b.read_boutframes_sheet(per_animal, 'CAB01_Alcohol')
    assert 'Session' not in df.columns


def test_a_session_tag_is_not_parsed_as_a_behavior():
    """Belt and braces: even handed the raw sheet, the parser ignores the tag."""
    df = pd.DataFrame({'Session': ['Alcohol'], 'Alcohol__start': [1], 'Alcohol__end': [2]})
    behaviors, _has_end = _G._parse_boutframes_dataframe(df)
    assert [n for n, _s, _e in behaviors] == ['Alcohol']


def test_session_matching_ignores_case_and_padding():
    b = _Bouts()
    df = pd.DataFrame({'Session': ['  alcohol '], 'A__start': [7]})
    assert list(b.filter_boutframes_rows(df, 'Alcohol')['A__start']) == [7]


def test_day_is_accepted_as_the_tag_column():
    b = _Bouts()
    df = pd.DataFrame({'Day': ['d1', 'd2'], 'A__start': [1, 2]})
    out = b.filter_boutframes_rows(df, 'd2')
    assert list(out['A__start']) == [2] and 'Day' not in out.columns


def test_a_session_with_no_scored_rows_yields_an_empty_frame(per_animal):
    b = _Bouts(['CAB03_Fentanyl'])
    df = b.read_boutframes_sheet(per_animal, 'CAB03_Fentanyl')
    assert df.empty
    behaviors, _ = _G._parse_boutframes_dataframe(df)
    assert all(len(s) == 0 for _n, s, _e in behaviors)


# ---------------------------------------------------------------------------
# The layouts that existed before
# ---------------------------------------------------------------------------

def test_per_recording_sheets_still_read_whole(per_recording):
    b = _Bouts(['CAB01_Alcohol', 'CAB01_Fentanyl'])
    assert b.resolve_boutframes_sheet(per_recording, 'CAB01_Alcohol') == ('CAB01_Alcohol', '')
    assert list(b.read_boutframes_sheet(per_recording, 'CAB01_Alcohol')['Alcohol__start']) == [156]


def test_a_single_session_project_is_untouched(single_session):
    b = _Bouts(['DG01'])
    assert b.resolve_boutframes_sheet(single_session, 'DG01') == ('DG01', '')
    df = b.read_boutframes_sheet(single_session, 'DG01')
    assert list(df['Drink__start']) == [10, 20]


def test_an_untagged_sheet_read_for_a_session_keeps_every_row(tmp_path):
    """A recording ID that splits, but a workbook that predates sessions."""
    path = tmp_path / 'bf.xlsx'
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        pd.DataFrame({'A__start': [1, 2]}).to_excel(w, sheet_name='CAB01', index=False)
    b = _Bouts(['CAB01_Alcohol'])
    assert list(b.read_boutframes_sheet(str(path), 'CAB01_Alcohol')['A__start']) == [1, 2]


# ---------------------------------------------------------------------------
# Workbook-wide views
# ---------------------------------------------------------------------------

def test_behavior_listing_excludes_the_session_tag(per_animal):
    b = _Bouts()
    assert 'Session' not in b.boutframes_behavior_columns(per_animal)
    assert 'Alcohol__start' in b.boutframes_behavior_columns(per_animal)


def test_iteration_yields_one_entry_per_recording_not_per_sheet(per_animal):
    b = _Bouts(['CAB01_Alcohol', 'CAB01_Fentanyl', 'CAB01_WaterSucrose', 'CAB03_Alcohol'])
    got = {label: df for label, _sid, df in b.iter_boutframes_recordings(per_animal)}
    assert set(got) == {'CAB01_Alcohol', 'CAB01_Fentanyl',
                        'CAB01_WaterSucrose', 'CAB03_Alcohol'}
    assert len(got['CAB01_Alcohol']) == 2
    assert len(got['CAB01_Fentanyl']) == 1


def test_iteration_still_splits_sessions_before_anything_is_processed(per_animal):
    """The exclusion preview runs with an empty project; it must not lump three
    days together, because the proximity rule would compare across days."""
    b = _Bouts()
    labels = [label for label, _sid, _df in b.iter_boutframes_recordings(per_animal)]
    assert sorted(labels) == ['CAB01 [Alcohol]', 'CAB01 [Fentanyl]',
                              'CAB01 [WaterSucrose]', 'CAB03 [Alcohol]']


def test_iteration_of_an_untagged_workbook_yields_the_sheets(single_session):
    b = _Bouts()
    assert [(l, s) for l, s, _ in b.iter_boutframes_recordings(single_session)] == [('DG01', None)]


def test_a_sheet_maps_back_to_the_recording_that_scored_a_behavior(per_animal):
    b = _Bouts(['CAB01_Alcohol', 'CAB01_Fentanyl', 'CAB01_WaterSucrose'])
    assert b.boutframes_recording_for_sheet(per_animal, 'CAB01', 'Fentanyl__start') == 'CAB01_Fentanyl'
    assert b.boutframes_recording_for_sheet(per_animal, 'CAB01', 'Water__start') == 'CAB01_WaterSucrose'


def test_sheet_to_recording_is_none_when_nothing_is_processed(per_animal):
    assert _Bouts().boutframes_recording_for_sheet(per_animal, 'CAB01') is None


# ---------------------------------------------------------------------------
# Writing back (random-shuffle / TTL merge)
# ---------------------------------------------------------------------------

def test_writing_one_session_leaves_the_others_intact(per_animal):
    """The bug this guards: writing per recording used to replace the whole
    sheet, so generating control bouts for one day erased the other two."""
    b = _Bouts(['CAB01_Alcohol', 'CAB01_Fentanyl'])
    wb = load_workbook(per_animal)
    header, blocks = b._read_sheet_columns_by_session(wb, 'CAB01', {'random__start'})
    assert header == 'Session'
    assert set(blocks) == {'Alcohol', 'Fentanyl', 'WaterSucrose'}

    blocks['Alcohol']['random__start'] = [11, 22, 33]
    b._rewrite_sheet_columns_by_session(wb, 'CAB01', header, blocks)
    wb.save(per_animal)

    alc = b.read_boutframes_sheet(per_animal, 'CAB01_Alcohol')
    fen = b.read_boutframes_sheet(per_animal, 'CAB01_Fentanyl')
    assert list(alc['random__start']) == [11, 22, 33]
    assert list(alc['Alcohol__start'].dropna()) == [156, 525]     # survived
    assert list(fen['Fentanyl__start'].dropna()) == [2567]        # survived
    assert 'random__start' not in fen.columns or fen['random__start'].isna().all()


def test_a_dropped_column_does_not_come_back(per_animal):
    b = _Bouts(['CAB01_Alcohol'])
    wb = load_workbook(per_animal)
    header, blocks = b._read_sheet_columns_by_session(wb, 'CAB01', {'alcohol__end'})
    assert all('Alcohol__end' not in block for block in blocks.values())


def test_an_untagged_sheet_reports_no_session_header(single_session):
    b = _Bouts(['DG01'])
    header, blocks = b._read_sheet_columns_by_session(load_workbook(single_session), 'DG01')
    assert header is None
    assert list(blocks) == ['']
    assert blocks['']['Drink__start'] == [10, 20]


def test_reading_a_missing_sheet_is_empty_not_an_error():
    b = _Bouts()
    wb = Workbook()
    assert b._read_sheet_columns_by_session(wb, 'nope') == (None, {})
